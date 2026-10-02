"""Quick end-to-end: login + GET /api/stats/export + verify xlsx contents."""
import base64
import os
import sys
import urllib.error
import urllib.request

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

BASE = "http://127.0.0.1:8000"
OUT = os.path.join(os.path.dirname(__file__), "..", "test_export_smoke.xlsx")

# 1) Login
auth = "Basic " + base64.b64encode(b"admin:admin123").decode()
req = urllib.request.Request(BASE + "/api/auth/login", method="POST",
                               headers={"Authorization": auth})
with urllib.request.urlopen(req, timeout=5) as r:
    cookie = (r.headers.get("Set-Cookie") or "").split(";")[0]
print(f"cookie: {cookie[:40]}...")

# 2) Export
req = urllib.request.Request(
    BASE + "/api/stats/export?month=9&year=2026",
    method="GET",
    headers={"Cookie": cookie},
)
try:
    with urllib.request.urlopen(req, timeout=15) as r:
        status = r.status
        blob = r.read()
        cd = r.headers.get("Content-Disposition", "")
        ct = r.headers.get("Content-Type", "")
except urllib.error.HTTPError as e:
    body = e.read()
    print(f"FAIL: {e.code} {e.reason}")
    print(body.decode("utf-8", errors="replace"))
    sys.exit(1)

print(f"status: {status}")
print(f"content-type: {ct}")
print(f"content-disposition: {cd}")
print(f"body size: {len(blob)} bytes")
if len(blob) < 1000:
    print("body too small:")
    print(blob[:500])
    sys.exit(1)

# 3) Save + inspect
with open(OUT, "wb") as f:
    f.write(blob)
print(f"saved {OUT} ({len(blob)} bytes)")

# 4) Verify the resulting xlsx
from openpyxl import load_workbook
import io
wb = load_workbook(io.BytesIO(blob))
print(f"sheets in output: {wb.sheetnames}")
if "T09_2026" not in wb.sheetnames:
    print(f"FAIL: expected sheet 'T09_2026', got {wb.sheetnames}")
    sys.exit(1)

ws = wb["T09_2026"]
print(f"T09_2026 dimensions: {ws.max_row} rows x {ws.max_column} cols")
print(f"A1: {ws['A1'].value!r}")
print(f"C1 (month): {ws['C1'].value}")
print(f"C2 (year): {ws['C2'].value}")
print(f"A5 (title): {ws['A5'].value!r}")
print(f"G9 (date): {ws['G9'].value}")
print(f"H9 formula: {ws['H9'].value!r}")
print(f"G11 weekday: {ws['G11'].value!r}")
# Spot-check body cells (should be empty or numeric after export cleared)
print(f"A12: {ws['A12'].value!r}")

# Cleanup
os.remove(OUT)
print("\nOK: end-to-end export works")